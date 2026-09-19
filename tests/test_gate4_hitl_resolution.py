"""Gate 4: Human-in-the-Loop Resolution & A2A Fleet Protocol Test."""

from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document

from app.integration.a2a_server import app


@pytest.fixture
def client():
    return TestClient(app)


def test_agent_card_contract(client):
    response = client.get("/.well-known/agent-card.json")
    assert response.status_code == 200
    card = response.json()
    assert card["id"] == "PPAC-REPORT-SYNTHESIS-AGENT"
    assert card["protocol"] == "a2a-jsonrpc-1.0"
    assert "generate_report_draft" in [m["name"] for m in card["methods"]]
    assert "resolve_audit_anomaly" in [m["name"] for m in card["methods"]]


def test_end_to_end_pipeline_and_hitl_resolution(client):
    # 1. Trigger autonomous scheduled run for August 2026
    trigger_resp = client.post("/api/v1/pipeline/trigger", json={"period_id": "2026-08"})
    assert trigger_resp.status_code == 200
    data = trigger_resp.json()

    assert data["status"] == "DRAFT_NEEDS_REVIEW"
    assert len(data["active_anomalies"]) >= 1

    bpcl_anomaly = [a for a in data["active_anomalies"] if a["entity_id"] == "BPCL"][0]
    anomaly_id = bpcl_anomaly["anomaly_id"]
    doc_uri = data["document_uri"]
    assert Path(doc_uri).exists()

    # 2. Check pipeline status endpoint
    status_resp = client.get("/api/v1/pipeline/status/2026-08")
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] == "DRAFT_NEEDS_REVIEW"

    # 3. Simulate Analyst resolving the audit flag programmatically
    resolve_resp = client.post(
        "/api/v1/anomalies/resolve",
        json={
            "anomaly_id": anomaly_id,
            "corrected_value": 445.0,
            "resolution_rationale": "Confirmed entry typo with BPCL coordinator; corrected to 445 TMT",
            "resolver_user": "lead.analyst@ppac.gov.in",
        },
    )
    assert resolve_resp.status_code == 200
    res_data = resolve_resp.json()
    assert res_data["status"] == "RESOLVED"
    assert res_data["document_status"] == "APPROVED_FOR_PUBLICATION"
    assert res_data["remaining_active_flags"] == 0

    # 4. Verify re-compiled document cleared Section 0 warnings
    recompiled_doc = Document(res_data["document_uri"])
    doc_text = "\n".join([p.text for p in recompiled_doc.paragraphs])
    table_text = "\n".join([c.text for t in recompiled_doc.tables for row in t.rows for c in row.cells])
    combined_text = doc_text + "\n" + table_text

    assert "ALL RECONCILIATION & ANOMALY CHECKS PASSED" in combined_text
    assert "445.0" in combined_text  # Updated value reflected in document tables


def test_a2a_jsonrpc_protocol_methods(client):
    rpc_payload = {
        "jsonrpc": "2.0",
        "method": "generate_report_draft",
        "params": {"period_id": "2026-08"},
        "id": 42,
    }
    resp = client.post("/a2a/ppac_report_synthesis/jsonrpc", json=rpc_payload)
    assert resp.status_code == 200
    res_body = resp.json()
    assert res_body["jsonrpc"] == "2.0"
    assert res_body["id"] == 42
    assert "document_uri" in res_body["result"]

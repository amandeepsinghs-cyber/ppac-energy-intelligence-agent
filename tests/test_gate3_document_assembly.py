"""Gate 3: Zero-to-Draft Document Assembly Test."""

from pathlib import Path
import pytest
from docx import Document

from app.canonical.models import ProductType, PricingBenchmark
from app.ingestion.excel_parser import OmcExcelParser
from app.ingestion.anomaly_engine import AnomalyEngine
from app.tsa.sarimax_model import FuelDemandForecaster
from app.tsa.chart_renderer import ForecastChartRenderer
from app.synthesis.docx_compiler import PpacDocxCompiler


@pytest.fixture
def fixtures_dir():
    return Path(__file__).parent.parent / "fixtures"


@pytest.fixture
def tmp_output_dir(tmp_path):
    return tmp_path / "output_docs"


def test_zero_to_draft_document_assembly(fixtures_dir, tmp_output_dir):
    # 1. Ingest Excel and flag anomalies
    excel_path = fixtures_dir / "mock_omc_sales_aug2026.xlsx"
    parser = OmcExcelParser()
    records, comparisons = parser.parse_workbook(excel_path, period_id="2026-08")

    anomaly_engine = AnomalyEngine()
    anomalies = anomaly_engine.evaluate_variances(comparisons, period_id="2026-08")

    # 2. Build pricing benchmark
    pricing = PricingBenchmark(
        period_id="2026-08",
        brent_usd_bbl=82.10,
        oman_dubai_usd_bbl=78.50,
        weight_sour_pct=75.6,
        weight_sweet_pct=24.4,
        icb_composite_usd=79.38,
        icb_composite_inr=round(79.38 * 83.95, 2),
        apm_natural_gas_usd_mmbtu=6.50,
        hpht_deepwater_gas_usd_mmbtu=9.87,
        usd_inr_exchange_rate=83.95,
    )

    # 3. Forecast and generate chart
    csv_path = fixtures_dir / "mock_historical_consumption_2018_2026.csv"
    forecaster = FuelDemandForecaster()
    forecast = forecaster.fit_and_forecast(csv_path, product=ProductType.HSD, forecast_steps=2)

    chart_renderer = ForecastChartRenderer()
    chart_path = tmp_output_dir / "test_hsd_chart.png"
    chart_renderer.render_forecast_chart(csv_path, forecast, chart_path)

    # 4. Compile Document
    compiler = PpacDocxCompiler()
    doc_out = tmp_output_dir / "Monthly_Flash_Report_Aug2026_Draft.docx"
    compiler.compile_flash_report(
        period_id="2026-08",
        consumption_records=records,
        pricing=pricing,
        anomalies=anomalies,
        forecast=forecast,
        output_path=doc_out,
        chart_image_path=chart_path,
    )

    assert doc_out.exists()
    assert doc_out.stat().st_size > 15000  # Multi-page docx with embedded chart

    # 5. Read back document and verify Section 0 & authentic PPAC content
    doc = Document(str(doc_out))
    p_text = "\n".join([p.text for p in doc.paragraphs])
    table_text = "\n".join([cell.text for t in doc.tables for row in t.rows for cell in row.cells])
    full_text = p_text + "\n" + table_text

    assert "SECTION 0: AI AUDIT & REVIEW CHECKLIST" in full_text
    assert "Chapter 1: Selected Macro-Economic Indicators" in full_text
    assert "Chapter 2: Indigenous Crude Production & Net Import Dependency" in full_text
    assert "87.8%" in full_text
    assert "Chapter 3: Refinery Crude Intake & Capacity Utilization" in full_text
    assert "Chapter 4: Domestic Petroleum Product Sales & OMC Breakdown" in full_text
    assert "Chapter 5: Indian Crude Basket (ICB) & International Prices" in full_text
    assert "Chapter 6: Natural Gas Administered Pricing & HPHT Deepwater Ceiling" in full_text
    assert "Chapter 7: Retail Selling Price (RSP) Build-up at Delhi" in full_text
    assert "Chapter 8: Econometric 60-Day Fuel Demand Forecast (SARIMAX)" in full_text
    assert "BPCL" in full_text
    assert "LPG" in full_text
    assert len(doc.tables) >= 5  # Section 0 callout + Macro + Indigenous + Sales + Pricing + RSP + Forecast
